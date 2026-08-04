#!/usr/bin/env python3
"""Package the exact HTML slide renders into a faithful 16:9 PowerPoint."""

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, "/tmp/nm-pptx-deps")

from pptx import Presentation
from pptx.util import Inches

OUT = ROOT / "Template_Slides_NeuralMind_AI_Fiel.pptx"
RENDERS = ROOT / "assets" / "slides-rendered"
TITLES = [
    "Capa", "Divisor de seção", "Verticais", "Produto com mockup",
    "Manifesto", "Título e lista", "Métrica", "Processo",
    "Encerramento", "Foto e statement", "Fundadores",
    "Encerramento com equipe", "Processo em três etapas", "Arquitetura de produto",
]

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

for number, title in enumerate(TITLES, 1):
    image = RENDERS / f"slide-{number:02d}.png"
    if not image.exists():
        raise SystemExit(f"Render ausente: {image}")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = (
        f"Template NeuralMind.AI — {title}.\n"
        "Este slide preserva com fidelidade a renderização do modelo HTML/CSS."
    )

prs.core_properties.title = "Template de Slides — NeuralMind.AI — versão fiel"
prs.core_properties.subject = "Renderização fiel dos 14 templates HTML/CSS em formato 16:9"
prs.core_properties.author = "NeuralMind.AI"
prs.core_properties.comments = "Slides renderizados diretamente do brand guide digital."
prs.save(OUT)
print(OUT)
