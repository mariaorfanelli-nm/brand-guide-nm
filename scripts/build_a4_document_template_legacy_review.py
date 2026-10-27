"""Build a review copy of the A4 template with the legacy NeuralMind logo."""

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

import build_a4_document_template as base


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "review" / "Template_A4_NeuralMind_logo_antigo.docx"

base.OUTPUT = OUTPUT
base.LOGO_PURPLE = ROOT / "assets" / "brand" / "legacy" / "neuralmind-legacy-color.png"
base.LOGO_WHITE = ROOT / "assets" / "brand" / "legacy" / "neuralmind-legacy-white.png"


def configure_legacy_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = base.Pt(8)
    shape = paragraph.add_run().add_picture(str(base.LOGO_PURPLE), width=base.Inches(1.55))
    base.set_image_alt(shape, "Logotipo legado neuralmind.ai")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = base.Pt(8)
    fp.paragraph_format.tab_stops.add_tab_stop(base.Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
    left = fp.add_run("neuralmind.ai  |  Documento institucional")
    base.set_run_font(left, size=8.5, color=base.MUTED)
    fp.add_run("\t")
    page_label = fp.add_run("PÁGINA ")
    base.set_run_font(page_label, size=8.5, color=base.PURPLE, bold=True)
    base.add_field(fp, "PAGE")


base.configure_header_footer = configure_legacy_header_footer


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    base.build()
