from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Template_A4_NeuralMind_AI.docx"
LOGO_PURPLE = ROOT / "assets" / "brand" / "neuralmind-ai-purple.png"
LOGO_WHITE = ROOT / "assets" / "brand" / "neuralmind-ai-white.png"

FONT = "Arial"
PURPLE = "290241"
VIOLET = "9381FF"
BLACK = "0F1108"
WHITE = "FFFFFC"
CORAL = "EF476F"
MUTED = "5C5861"
LAVENDER = "F3F0FF"
HAIRLINE = "D8D1DE"

def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color, bold=False, italic=False):
    style.font.name = FONT
    r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_border(paragraph, side, color, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    border = borders.find(qn(f"w:{side}"))
    if border is None:
        border = OxmlElement(f"w:{side}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=140):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def set_table_borders(table, color=HAIRLINE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
    for edge in ("left", "right", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_image_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_spacer(doc, points):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(points)
    paragraph.paragraph_format.line_spacing = Pt(1)
    return paragraph


def add_styled_paragraph(doc, text, style=None, size=None, color=BLACK, bold=None,
                         italic=None, before=0, after=6, alignment=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if alignment is not None:
        paragraph.alignment = alignment
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_callout(doc, label, text, fill=LAVENDER, accent=VIOLET):
    paragraph = doc.add_paragraph(style="NM Callout")
    set_paragraph_shading(paragraph, fill)
    set_paragraph_border(paragraph, "top", accent, size=12, space=7)
    label_run = paragraph.add_run(label.upper() + "\n")
    set_run_font(label_run, size=8, color=PURPLE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11.5, color=PURPLE, bold=False)
    return paragraph


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 10.5, BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_style_font(title, 28, PURPLE, bold=True)
    title.paragraph_format.space_after = Pt(7)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, 12.5, MUTED, italic=False)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 17.5, PURPLE, 20, 7),
        ("Heading 2", 12, BLACK, 13, 5),
        ("Heading 3", 10.5, PURPLE, 10, 4),
    ):
        style = styles[style_name]
        set_style_font(style, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        set_style_font(style, 10.5, BLACK)
        style.paragraph_format.left_indent = Cm(0.8)
        style.paragraph_format.first_line_indent = Cm(-0.45)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.25

    custom = {
        "NM Eyebrow": (WD_STYLE_TYPE.PARAGRAPH, 8, VIOLET, True),
        "NM Lead": (WD_STYLE_TYPE.PARAGRAPH, 12.5, PURPLE, False),
        "NM Callout": (WD_STYLE_TYPE.PARAGRAPH, 10.5, PURPLE, False),
        "NM Caption": (WD_STYLE_TYPE.PARAGRAPH, 8.5, MUTED, False),
        "NM Quote": (WD_STYLE_TYPE.PARAGRAPH, 13.5, PURPLE, False),
    }
    for name, (style_type, size, color, bold) in custom.items():
        if name not in styles:
            style = styles.add_style(name, style_type)
        else:
            style = styles[name]
        set_style_font(style, size, color, bold=bold)
        style.paragraph_format.widow_control = True

    styles["NM Eyebrow"].paragraph_format.space_after = Pt(6)
    styles["NM Lead"].paragraph_format.space_after = Pt(15)
    styles["NM Lead"].paragraph_format.line_spacing = 1.34
    styles["NM Callout"].paragraph_format.space_before = Pt(8)
    styles["NM Callout"].paragraph_format.space_after = Pt(12)
    styles["NM Callout"].paragraph_format.line_spacing = 1.3
    styles["NM Caption"].paragraph_format.space_before = Pt(4)
    styles["NM Caption"].paragraph_format.space_after = Pt(8)
    styles["NM Quote"].paragraph_format.space_before = Pt(10)
    styles["NM Quote"].paragraph_format.space_after = Pt(10)
    styles["NM Quote"].paragraph_format.line_spacing = 1.25
    styles["NM Quote"].paragraph_format.keep_together = True


def configure_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.15)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = False


def configure_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run().add_picture(str(LOGO_PURPLE), width=Inches(1.8))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(8)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT)
    left = fp.add_run("NeuralMind.AI  |  Documento institucional")
    set_run_font(left, size=8.5, color=MUTED)
    fp.add_run("\t")
    page_label = fp.add_run("PÁGINA ")
    set_run_font(page_label, size=8.5, color=PURPLE, bold=True)
    add_field(fp, "PAGE")


def build_cover(doc):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.columns[0].width = Cm(21.0)
    row = table.rows[0]
    row.height = Cm(29.68)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    set_repeat_table_header(row)
    cell = row.cells[0]
    cell.width = Cm(21.0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(cell, PURPLE)
    set_cell_margins(cell, top=900, start=1815, bottom=800, end=1815)

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    tbl_pr.append(borders)

    logo_p = cell.paragraphs[0]
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.paragraph_format.space_after = Pt(118)
    shape = logo_p.add_run().add_picture(str(LOGO_WHITE), width=Inches(2.35))
    set_image_alt(shape, "Logotipo NeuralMind.AI")

    eyebrow = cell.add_paragraph(style="NM Eyebrow")
    eyebrow.paragraph_format.space_after = Pt(10)
    eyebrow_run = eyebrow.add_run("Documento institucional  /  modelo A4")
    set_run_font(eyebrow_run, color=VIOLET, bold=True)

    title = cell.add_paragraph(style="Title")
    title_run = title.add_run("[Título do documento]")
    set_run_font(title_run, size=28, color=WHITE, bold=True)
    subtitle = cell.add_paragraph(style="Subtitle")
    subtitle_run = subtitle.add_run("[Uma frase curta que explica o propósito desta publicação]")
    set_run_font(subtitle_run, size=12.5, color=WHITE)

    spacer = cell.add_paragraph()
    spacer.paragraph_format.space_after = Pt(126)
    spacer.add_run("")

    prepared = cell.add_paragraph(style="NM Eyebrow")
    prepared.paragraph_format.space_after = Pt(1)
    set_run_font(prepared.add_run("Preparado por"), color=VIOLET, bold=True)
    owner = cell.add_paragraph()
    owner.paragraph_format.space_after = Pt(11)
    set_run_font(owner.add_run("[Área ou responsável]"), size=10.5, color=WHITE)
    date_label = cell.add_paragraph(style="NM Eyebrow")
    date_label.paragraph_format.space_after = Pt(1)
    set_run_font(date_label.add_run("Data  /  versão"), color=VIOLET, bold=True)
    date_value = cell.add_paragraph()
    date_value.paragraph_format.space_after = Pt(0)
    set_run_font(date_value.add_run("[dd.mm.aaaa]  /  [v1.0]"), size=10.5, color=WHITE)


def add_body_template(doc, usable_width_dxa):
    note = doc.add_paragraph(style="NM Caption")
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(16)
    note_run = note.add_run("Orientação do modelo: substitua os textos entre colchetes e exclua esta linha antes de publicar.")
    set_run_font(note_run, size=8.5, color=MUTED, italic=True)

    doc.add_heading("Resumo executivo", level=1)
    lead = doc.add_paragraph(style="NM Lead")
    lead.add_run("[Apresente o objetivo do documento, a conclusão principal e o impacto esperado em um parágrafo curto e direto.]")

    key = doc.add_paragraph(style="NM Callout")
    key.paragraph_format.space_before = Pt(8)
    key.paragraph_format.space_after = Pt(18)
    label = key.add_run("Em uma frase\n")
    set_run_font(label, size=8, color=VIOLET, bold=True)
    key_text = key.add_run("[Registre a ideia que precisa permanecer na memória do leitor.]")
    set_run_font(key_text, size=11, color=PURPLE)

    doc.add_heading("Contexto", level=1)
    add_styled_paragraph(
        doc,
        "[Explique o cenário com fatos, referências e vocabulário técnico preciso. Dê ao leitor apenas o contexto necessário para compreender a análise.]",
    )
    add_styled_paragraph(
        doc,
        "[Conecte o problema às pessoas, aos processos e às decisões envolvidas. Prefira frases naturais e concretas a fórmulas institucionais genéricas.]",
    )

    doc.add_heading("Objetivos", level=2)
    for text in (
        "[O que este trabalho pretende alcançar.]",
        "[Como o resultado será reconhecido ou medido.]",
        "[Quais limites ou dependências devem ser considerados.]",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

    doc.add_heading("Plano de ação", level=1)
    add_styled_paragraph(
        doc,
        "[Descreva a execução como uma sequência real. Cada etapa deve indicar uma ação, seu responsável e o resultado esperado.]",
    )

    for text in (
        "[Alinhar o escopo e os critérios de sucesso.]",
        "[Executar o trabalho com rastreabilidade e governança.]",
        "[Revisar os resultados e decidir a próxima ação.]",
    ):
        p = doc.add_paragraph(style="List Number")
        p.add_run(text)

    data_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(data_section)
    configure_header_footer(data_section)
    doc.add_heading("Indicadores e acompanhamento", level=2)
    caption = doc.add_paragraph(style="NM Caption")
    caption.add_run("Exemplo de estrutura para acompanhamento")

    table = doc.add_table(rows=1, cols=4)
    widths = [2350, 3460, 1700, usable_width_dxa - 7510]
    set_table_geometry(table, widths, indent_dxa=140)
    set_table_borders(table, color="C9C1D1", size=6)
    headers = ["Indicador", "Definição", "Meta", "Situação"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, WHITE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=8.5, color=PURPLE, bold=True)
    set_repeat_table_header(table.rows[0])

    data = [
        ("[Indicador]", "[O que será medido e por quê.]", "[Meta]", "[Em curso]"),
        ("[Indicador]", "[Fonte, frequência e responsável.]", "[Meta]", "[Planejado]"),
        ("[Indicador]", "[Critério objetivo de conclusão.]", "[Meta]", "[Concluído]"),
    ]
    for row_index, values in enumerate(data, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, LAVENDER)
        for column, value in enumerate(values):
            cell = cells[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if column >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, size=9, color=BLACK, bold=(column == 0))
    set_table_geometry(table, widths, indent_dxa=140)

    doc.add_heading("Ponto de atenção", level=3)
    attention = doc.add_paragraph()
    attention.paragraph_format.space_before = Pt(4)
    attention.paragraph_format.space_after = Pt(16)
    attention.add_run("[Registre riscos, restrições, decisões pendentes ou requisitos de governança com linguagem direta.]")

    doc.add_heading("Conclusão", level=1)
    add_styled_paragraph(
        doc,
        "[Retome a mensagem principal, explicite a decisão esperada e indique o próximo marco. Termine com uma frase direta e acionável.]",
    )
    doc.add_heading("Próximas ações", level=2)
    for text in (
        "[Ação imediata / responsável / prazo]",
        "[Validação necessária / responsável / prazo]",
        "[Próximo ponto de controle / data]",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

def set_document_properties(doc):
    props = doc.core_properties
    props.title = "Template A4 NeuralMind.AI"
    props.subject = "Template institucional de documentos A4"
    props.author = "NeuralMind.AI"
    props.keywords = "NeuralMind.AI, template, A4, marca, documento institucional"
    props.comments = "Template criado conforme o Brand Guide NeuralMind.AI."


def build():
    doc = Document()
    configure_styles(doc)
    cover_section = doc.sections[0]
    cover_section.page_width = Cm(21.0)
    cover_section.page_height = Cm(29.7)
    cover_section.top_margin = Cm(0)
    cover_section.bottom_margin = Cm(0)
    cover_section.left_margin = Cm(0)
    cover_section.right_margin = Cm(0)
    cover_section.header_distance = Cm(0)
    cover_section.footer_distance = Cm(0)
    set_document_properties(doc)

    build_cover(doc)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body_section)
    configure_header_footer(body_section)
    usable_width_dxa = round(
        (body_section.page_width - body_section.left_margin - body_section.right_margin) / 635
    )
    add_body_template(doc, usable_width_dxa)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
