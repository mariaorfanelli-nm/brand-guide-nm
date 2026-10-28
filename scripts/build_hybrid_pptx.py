#!/usr/bin/env python3
"""Build a faithful PowerPoint with rendered backgrounds and editable text boxes."""

from pathlib import Path
import json
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, "/tmp/nm-pptx-deps")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

ASSETS = ROOT / "assets" / "slides-hybrid"
OUT = ROOT / "Template_Slides_NeuralMind_AI_Editavel.pptx"
PX_PER_INCH = 120.0


def parse_color(value):
    nums = [int(v) for v in re.findall(r"\d+", value)[:3]]
    return RGBColor(*nums)


def alignment(value):
    if value in ("center",): return PP_ALIGN.CENTER
    if value in ("right", "end"): return PP_ALIGN.RIGHT
    if value == "justify": return PP_ALIGN.JUSTIFY
    return PP_ALIGN.LEFT


def add_text(slide, item):
    value = item["text"]
    if item["textTransform"] == "uppercase": value = value.upper()
    x = item["x"] / PX_PER_INCH
    y = item["y"] / PX_PER_INCH
    w = max(item["w"] / PX_PER_INCH + 0.025, 0.08)
    h = max(item["h"] / PX_PER_INCH + 0.04, 0.08)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.auto_size = MSO_AUTO_SIZE.NONE; frame.word_wrap = False; frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    lines = value.split("\n")
    for idx, line in enumerate(lines):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.alignment = alignment(item["textAlign"])
        p.space_before = p.space_after = Pt(0)
        line_height = item.get("lineHeight", "normal")
        if isinstance(line_height, str) and line_height.endswith("px"):
            p.line_spacing = Pt(float(line_height[:-2]) * 0.6)
        run = p.add_run(); run.text = line
        font = run.font
        family = "Urbanist" if "Urbanist" in item.get("fontFamily", "") else "Outfit"
        font.name = family
        font.size = Pt(float(item["fontSize"]) * 0.6)
        try: weight = int(item["fontWeight"])
        except ValueError: weight = 700 if item["fontWeight"] == "bold" else 400
        variant = {500: "Medium", 600: "SemiBold", 800: "ExtraBold", 900: "Black"}.get(weight)
        if variant: font.name = f"{family} {variant}"
        font.bold = weight == 700
        font.color.rgb = parse_color(item["color"])
    return box


prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

for number in range(1, 17):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = ASSETS / f"slide-{number:02d}-background.png"
    metadata = json.loads((ASSETS / f"slide-{number:02d}.json").read_text(encoding="utf-8"))
    slide.shapes.add_picture(str(background), 0, 0, width=prs.slide_width, height=prs.slide_height)
    for item in metadata:
        lines = item.get("renderedLines")
        if lines:
            # Use the browser's actual line breaks so Office cannot reflow titles.
            for index, line in enumerate(lines):
                y = item["y"] + index * (float(item["lineHeight"][:-2]) if item["lineHeight"].endswith("px") else item["fontSize"] * 1.2)
                part = {**item, "text": line["text"], "x": line["x"], "y": y, "w": max(line["w"] * 1.4 + 24, 40), "h": max(item["fontSize"] * 1.5, line["h"]), "textAlign": "left"}
                add_text(slide, part)
        else:
            add_text(slide, item)
    slide.notes_slide.notes_text_frame.text = (
        "Template NeuralMind.ai com fundo fiel ao HTML/CSS e textos editáveis.\n"
        "Para máxima fidelidade no PowerPoint, instale as famílias Outfit e Urbanist."
    )

prs.core_properties.title = "Template de Slides — NeuralMind.ai — editável"
prs.core_properties.subject = "Fundos fiéis ao HTML/CSS com caixas de texto editáveis"
prs.core_properties.author = "NeuralMind.ai"
prs.save(OUT)
print(OUT)
