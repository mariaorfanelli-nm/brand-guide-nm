#!/usr/bin/env python3
"""Build the editable NeuralMind.AI 16:9 PowerPoint template."""

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, "/tmp/nm-pptx-deps")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = ROOT / "Template_Slides_NeuralMind_AI.pptx"
ASSETS = ROOT / "assets"

W, H = 13.333333, 7.5
PURPLE = "290241"
PURPLE_2 = "3D0666"
INK = "0F1108"
CREAM = "F3EFE5"
WHITE = "FFFFFC"
CORAL = "F2766B"
LAVENDER = "9381FF"
MUTED = "6A6570"
FONT = "NeuralMind Sans Square"


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(0.7)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
         margin=0, spacing=1.0, caps=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    run = p.add_run(); run.text = value.upper() if caps else value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return box


def rich_text(slide, parts, x, y, w, h, size=20, color=INK, align=PP_ALIGN.LEFT,
              valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = box.text_frame; f.clear(); f.word_wrap = True
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    f.vertical_anchor = valign
    p = f.paragraphs[0]; p.alignment = align
    for value, kwargs in parts:
        r = p.add_run(); r.text = value; r.font.name = FONT; r.font.size = Pt(kwargs.get("size", size))
        r.font.bold = kwargs.get("bold", False); r.font.color.rgb = rgb(kwargs.get("color", color))
    return box


def picture(slide, rel, x, y, w=None, h=None):
    kwargs = {}
    if w is not None: kwargs["width"] = Inches(w)
    if h is not None: kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(ASSETS / rel), Inches(x), Inches(y), **kwargs)


def crop_picture(slide, rel, x, y, w, h):
    pic = picture(slide, rel, x, y, w, h)
    # python-pptx stretches when both dimensions are given; acceptable for these editorial crops.
    return pic


def bg(slide, color):
    rect(slide, 0, 0, W, H, color)


def watermark(slide, dark=False):
    rel = "brand/neuralmind-ai-icon-white.png" if dark else "brand/neuralmind-ai-icon-purple.png"
    pic = picture(slide, rel, 12.15, 6.82, 0.72)
    pic.image
    return pic


def eyebrow(slide, value, x, y, color=PURPLE):
    return text(slide, value, x, y, 5.6, 0.28, 10, color, True, caps=True)


def add_notes(slide, title):
    slide.notes_slide.notes_text_frame.text = (
        f"Template NeuralMind.AI — {title}.\n"
        "Substitua os textos de exemplo mantendo hierarquia, margens e contraste."
    )


def add_slide(prs, title, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, PURPLE if dark else CREAM)
    add_notes(slide, title)
    return slide


def cover(prs):
    s = add_slide(prs, "Capa", True)
    rect(s, 0, 0, W, H, "210235")
    rect(s, 0, 0, 4.0, 7.5, PURPLE_2, transparency=20)
    eyebrow(s, "Brand Guide Digital", 1.05, 4.18, LAVENDER)
    picture(s, "brand/neuralmind-ai-white.png", 1.05, 4.68, 5.25)
    rect(s, 1.05, 5.69, 1.05, 0.045, CORAL)
    text(s, "IA responsável para setores críticos.", 1.05, 5.96, 5.8, 0.48, 18, WHITE)


def section(prs):
    s = add_slide(prs, "Divisor de seção", True)
    eyebrow(s, "Capítulo 02", .95, 2.05, LAVENDER)
    text(s, "Arquitetura de\nProdutos", .95, 2.47, 9.5, 1.65, 45, WHITE, True, spacing=.92)
    text(s, "Três verticais, um único compromisso: IA responsável para decisões em ambientes regulados.",
         .95, 4.52, 6.8, .72, 17, WHITE)
    watermark(s, True)


def verticals(prs):
    s = add_slide(prs, "Visão geral de verticais", True)
    eyebrow(s, "Verticais de atuação", .82, .62, CORAL)
    text(s, "IA responsável para setores\nonde precisão não é opcional.", .82, 1.08, 6.0, 1.1, 29, WHITE, True)
    text(s, "Cada vertical resolve um tipo distinto de problema crítico, sempre com fonte verificável, rastreabilidade e supervisão humana no centro da decisão.", 7.2, 1.17, 5.2, .9, 13, WHITE)
    cards = [
        ("01", "Dados & Conhecimento", "Transforma documentos, bases internas e informação não estruturada em respostas, evidências e insights verificáveis.", "NeuralSearchX   NDocs   Tuss-Finder"),
        ("02", "Jurídico", "Modelos e aplicações especializados no contexto brasileiro para apoiar análise, fundamentação e produção jurídica.", "Jurema   ARACI   AssertIA   SofIA"),
        ("03", "Saúde", "Soluções para priorização, classificação e apoio operacional em ambientes clínicos, regulatórios e assistenciais.", "PrioScan"),
    ]
    for i, (n, t, body, products) in enumerate(cards):
        x = .82 + i * 4.18
        rect(s, x, 3.02, 3.82, 3.55, "3A1550", "725D80", True)
        text(s, n, x+.22, 3.25, .6, .22, 9, CORAL, True)
        text(s, t, x+.22, 3.66, 3.35, .6, 18, WHITE, True)
        text(s, body, x+.22, 4.38, 3.33, 1.02, 10.5, WHITE)
        rect(s, x+.22, 5.57, 3.34, .01, "725D80")
        text(s, products, x+.22, 5.78, 3.3, .45, 8.5, LAVENDER, True)


def product(prs):
    s = add_slide(prs, "Produto com mockup")
    eyebrow(s, "Produto exemplo", .78, .62)
    text(s, "NeuralSearchX", .78, 1.05, 5.1, .62, 34, PURPLE, True)
    text(s, "IA que responde o que você precisa — com a fonte sempre à vista.", .78, 1.82, 4.4, .82, 18, PURPLE, True)
    text(s, "Busca conversacional sobre bases internas de documentos. Cada resposta carrega o trecho original que a fundamenta, permitindo verificação imediata e decisões com lastro.", .78, 2.83, 4.35, 1.05, 11.5, MUTED)
    for i, item in enumerate(["Busca semântica multifonte", "Citação do trecho original", "Validação humana no fluxo"]):
        y = 4.18 + i*.64; rect(s, .78, y, 4.45, .01, "D4CCDA")
        text(s, f"0{i+1}", .78, y+.18, .45, .23, 8, CORAL, True)
        text(s, item, 1.42, y+.13, 3.7, .32, 10.5, INK)
    rect(s, 5.72, .55, 6.83, 6.35, "FAF9F4", "D4CCDA", True)
    rect(s, 6.13, 1.05, 6.0, 5.35, WHITE, "C9C0D0", True)
    for i in range(3): rect(s, 6.42+i*.18, 1.28, .08, .08, "D4CCDA", radius=True)
    eyebrow(s, "Consulta", 6.45, 1.77)
    text(s, "Quais contratos exigem revisão regulatória?", 6.45, 2.06, 5.25, .5, 13, INK, True)
    rect(s, 6.45, 2.86, 5.35, 1.35, "F3F0F7", "D4CCDA", True)
    eyebrow(s, "Resposta com fonte", 6.68, 3.1)
    text(s, "Foram identificados três documentos com cláusulas de atualização regulatória e prazo de revisão em aberto.", 6.68, 3.48, 4.75, .55, 10.5, INK)
    for i, val in enumerate(["Contrato matriz · pág. 14", "Anexo jurídico · pág. 08"]):
        y=4.55+i*.68; rect(s, 6.45, y, 5.35, .52, "FAF8FC", "DED7E3", True)
        text(s, f"Fonte 0{i+1}", 6.67, y+.1, 1.0, .2, 7.5, CORAL, True)
        text(s, val, 8.0, y+.08, 3.4, .24, 9.5, PURPLE, True)
    watermark(s)


def statement(prs):
    s = add_slide(prs, "Frase manifesto")
    text(s, "A IA gera, organiza, sugere e acelera.\nO humano valida, decide e assume o controle.",
         1.25, 2.12, 10.83, 2.0, 33, PURPLE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, "MANIFESTO NEURALMIND.AI", 4.65, 4.68, 4.0, .28, 9, CORAL, True, PP_ALIGN.CENTER)
    watermark(s)


def title_list(prs):
    s = add_slide(prs, "Título e lista")
    eyebrow(s, "Posicionamento", .95, 1.05)
    text(s, "Inteligência artificial\ncom rastreabilidade.", .95, 1.58, 5.3, 1.45, 34, PURPLE, True)
    items = ["Cada resposta carrega o trecho exato que a fundamenta.", "Modelos especializados em domínio — não generalistas.", "Operação auditável e em conformidade com a LGPD."]
    for i, item in enumerate(items):
        y=1.28+i*1.45; rect(s, 7.1, y, 5.2, .01, "CFC6D4")
        text(s, f"0{i+1}", 7.1, y+.35, .55, .3, 10, CORAL, True)
        text(s, item, 8.0, y+.27, 4.15, .65, 15, INK)
    watermark(s)


def metric(prs):
    s = add_slide(prs, "Métrica de impacto", True)
    eyebrow(s, "Resultado · NeuralSearchX 2.0", .95, 1.1, CORAL)
    rich_text(s, [("4", {"size":112, "bold":True, "color":WHITE}), ("×", {"size":66, "bold":True, "color":CORAL})], .9, 1.72, 5.2, 2.0)
    text(s, "Redução de latência da busca semântica multifonte na nova arquitetura, sem perda de citação ou rastreabilidade.", 7.3, 2.35, 4.8, 1.25, 19, WHITE)
    watermark(s, True)


def process(prs):
    s = add_slide(prs, "Fluxo de processo")
    eyebrow(s, "Fluxo de trabalho", .78, .62)
    text(s, "Da pergunta à decisão auditável.", .78, 1.03, 7.0, .58, 30, PURPLE, True)
    steps = [("Pergunta natural", "O usuário descreve a necessidade em linguagem comum."), ("Busca semântica", "Recuperação em múltiplas fontes com citação do trecho exato."), ("Síntese contextual", "Resposta consolidada, vinculada a cada referência."), ("Validação humana", "O especialista revisa, decide e assume o controle.")]
    rect(s, 1.18, 3.18, 10.65, .035, "CFC6D4")
    for i,(title,body) in enumerate(steps):
        x=.78+i*3.08
        rect(s, x, 2.86, .66, .66, PURPLE, radius=True)
        text(s, f"0{i+1}", x, 3.02, .66, .2, 9, WHITE, True, PP_ALIGN.CENTER)
        text(s, title, x, 3.82, 2.62, .42, 16, PURPLE, True)
        text(s, body, x, 4.42, 2.62, .9, 10.5, MUTED)
    watermark(s)


def closing(prs):
    s = add_slide(prs, "Encerramento", True)
    eyebrow(s, "Próximo passo", .95, 1.12, CORAL)
    text(s, "Vamos construir\njuntos.", .95, 1.62, 6.0, 1.5, 42, WHITE, True)
    picture(s, "brand/neuralmind-ai-white.png", .95, 5.62, 3.15)
    contacts=[("Comercial","contato@neuralmind.ai"),("Web","neuralmind.ai"),("LinkedIn","@neuralmind.ai")]
    for i,(label,val) in enumerate(contacts):
        y=1.43+i*1.45; eyebrow(s,label,8.15,y,LAVENDER); text(s,val,8.15,y+.38,3.9,.35,15,WHITE)


def photo_split(prs):
    s = add_slide(prs, "Foto e statement", True)
    crop_picture(s, "photos/collaboration.jpg", 0, 0, 6.2, 7.5)
    rect(s, 0, 0, 6.2, 7.5, PURPLE, transparency=68)
    rect(s, 6.2, 0, 7.14, 7.5, PURPLE)
    eyebrow(s, "Pessoas", 6.95, 1.45, CORAL)
    text(s, "A IA é poderosa.\nAs pessoas,\nindispensáveis.", 6.95, 1.95, 5.6, 2.1, 34, WHITE, True)
    text(s, "Construímos modelos para apoiar especialistas — não para substituí-los. A última palavra é, e segue sendo, humana.", 6.95, 4.58, 4.95, 1.05, 14, WHITE)


def founders(prs):
    s=add_slide(prs,"Fundadores")
    eyebrow(s,"Quem somos",.72,.58)
    text(s,"Ciência aplicada.\nResponsabilidade real.",.72,1.02,4.65,1.18,29,PURPLE,True)
    text(s,"Fundada por Patricia Magalhães e Roberto Lotufo, a NeuralMind transforma pesquisa em IA responsável, auditável e confiável para organizações que lidam com alta complexidade e decisões críticas.",.72,2.5,4.5,1.28,11.5,MUTED)
    for i,val in enumerate(["Rigor científico","Julgamento humano","Setores regulados"]):
        rect(s,.72,4.22+i*.5,3.85,.34,"E7E0EA",radius=True); text(s,val,.88,4.28+i*.5,3.4,.16,8.5,PURPLE,True)
    pics=[("photos/patricia-founder.jpg","Patricia Magalhães","Co-fundadora & CEO"),("photos/roberto-lotufo-founder.jpg","Roberto Lotufo","Co-fundador & CTO")]
    for i,(rel,name,role) in enumerate(pics):
        x=5.72+i*3.35; crop_picture(s,rel,x,.78,3.0,4.62); rect(s,x,4.46,3.0,.94,PURPLE,transparency=6)
        text(s,name,x+.2,4.68,2.6,.25,12,WHITE,True); text(s,role,x+.2,5.04,2.6,.2,8.5,LAVENDER)
    watermark(s)


def closing_photo(prs):
    s=add_slide(prs,"Encerramento com equipe",True)
    crop_picture(s,"photos/team-neuralmind-slide.jpg",0,0,W,H)
    rect(s,0,0,W,H,PURPLE,transparency=28)
    eyebrow(s,"Próximo passo",.85,.82,CORAL)
    text(s,"Vamos construir juntos.",.85,1.25,7.8,.72,38,WHITE,True)
    picture(s,"brand/neuralmind-ai-white.png",.85,6.22,3.1)
    eyebrow(s,"Comercial",8.25,6.12,LAVENDER); text(s,"contato@neuralmind.ai",8.25,6.43,2.4,.28,10,WHITE)
    eyebrow(s,"Web",10.78,6.12,LAVENDER); text(s,"neuralmind.ai",10.78,6.43,1.7,.28,10,WHITE)


def three_step(prs):
    s=add_slide(prs,"Processo em três colunas")
    eyebrow(s,"Processo em 3 etapas",.78,.58)
    text(s,"Da complexidade à decisão assistida.",.78,.98,7.2,.55,29,PURPLE,True)
    text(s,"Modelo para explicar jornadas curtas, pipelines comerciais, onboarding ou etapas de implantação.",.78,1.64,7.5,.45,11,MUTED)
    cards=[("01","Mapear","Identificamos fontes, regras de negócio, usuários e pontos de decisão que exigem rastreabilidade."),("02","Construir","Combinamos modelos especializados, busca semântica e agentes para criar uma solução sob medida."),("03","Validar","Especialistas revisam saídas, monitoram qualidade e transformam IA em operação confiável.")]
    for i,(n,t,b) in enumerate(cards):
        x=.78+i*4.18; rect(s,x,2.64,3.78,3.62,WHITE,"D5CCD9",True)
        text(s,n,x+.28,2.96,.55,.26,10,CORAL,True); text(s,t,x+.28,3.48,3.1,.45,20,PURPLE,True)
        text(s,b,x+.28,4.25,3.15,1.2,11,MUTED)
    watermark(s)


def product_stack(prs):
    s=add_slide(prs,"Arquitetura de produto",True)
    eyebrow(s,"Arquitetura de produto",.78,.55,CORAL)
    text(s,"ARACI",.78,.96,3.4,.55,31,WHITE,True)
    text(s,"Copiloto jurídico powered by NeuralSearchX + Jurema.",.78,1.62,6.0,.38,14,LAVENDER,True)
    rect(s,4.1,2.18,5.15,1.43,"4A1A61","80628F",True)
    eyebrow(s,"Produto principal",4.38,2.45,LAVENDER); text(s,"ARACI",4.38,2.83,1.35,.38,21,WHITE,True)
    text(s,"Análise, fundamentos e minutas com rastreabilidade para o trabalho jurídico.",5.9,2.7,2.95,.55,10.5,WHITE)
    blocks=[("Busca proprietária","NeuralSearchX","Recupera evidências, documentos e fontes verificáveis para cada resposta."),("Modelo especializado","Jurema","LLM jurídico brasileiro para linguagem, doutrina e contexto normativo local."),("Camada operacional","Agentes & Workflow","Orquestra tarefas, validações e próximos passos sob supervisão humana.")]
    for i,(lab,t,b) in enumerate(blocks):
        x=.78+i*4.18; rect(s,x,4.25,3.78,2.25,"371148","725D80",True)
        eyebrow(s,lab,x+.24,4.54,LAVENDER); text(s,t,x+.24,4.93,3.25,.4,17,WHITE,True); text(s,b,x+.24,5.5,3.2,.62,9.5,WHITE)


def build():
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
    for fn in [cover,section,verticals,product,statement,title_list,metric,process,closing,photo_split,founders,closing_photo,three_step,product_stack]: fn(prs)
    prs.core_properties.title="Template de Slides — NeuralMind.AI"
    prs.core_properties.subject="Sistema modular 16:9 editável"
    prs.core_properties.author="NeuralMind.AI"
    prs.core_properties.comments="Gerado a partir do brand guide digital em HTML/CSS."
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
